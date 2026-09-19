#!/usr/bin/env python3
"""Build the final results spreadsheet (.xlsx) from the runs in output-final/.

One tab per experiment, laid out like the "Resultados SSP-USPrC" Google Sheet
(blocks PT / Practitioner / MODELO / GA for Same Toolset, Different Toolset
and Dang reproduzidas; one row per instance for the Beezao IPMTC II tabs),
plus:
  LEIA-ME      where each tab comes from and how each column is computed
  Validação    validation errors/warnings per row (uteis/validador.py)
  Correções    cells where the recomputed value differs from the one the solver
               printed (only with --values recomputed)

Every report is validated first. --values picks where the counters and the
objective come from:
  recomputed   recomputed from the printed schedule by the validator (default)
  reported     as printed in the report footer -- what the original sheet used

--compare DIR checks the generated values cell by cell against the Google
Sheet tabs exported as CSV ("Resultados SSP-USPrC - <tab>.csv" in DIR), the
way output-final/ was matched to the sheet.

Usage:
  python3 scripts/buildSpreadsheet.py
  python3 scripts/buildSpreadsheet.py --values reported --compare .tmp --no-write
"""
import argparse
import csv
import datetime
import os
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from uteis import results as rs  # noqa: E402
from uteis import validador as vd  # noqa: E402

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SHEET_TAB = {"IPMTC II Practitioner": "IPMTC II Pratictitioners e GA"}  # sheet tab names that differ

# PT parameters as written in the original sheet's INFO column (not recorded in
# the reports). The base runs' PTL is 500 -- see uteis/results.py TABS.
PT_PARAMS = {
    "Same Toolset": "--TEMP_INIT 0.1 --TEMP_FIM 5 --N_REPLICAS 11 --MCL 500 --PTL 600 --PASSO_GATILHO 10 "
                    "--TEMP_DIST 3 --TYPE_UPDATE 1 --INIT_SOL_TYPE 0 --TEMP_UPDATE 3500 "
                    "--PTL_TEMP_UPDATE_PROPORTION 3",
    "Different Toolset": "--TEMP_INIT 0.1 --TEMP_FIM 10 --N_REPLICAS 11 --MCL 500 --PTL 600 --TEMP_DIST 1 "
                         "--TYPE_UPDATE 1 --INIT_SOL_TYPE 1 --PTL_TEMP_UPDATE_PROPORTION 5",
    "Dang reproduzidas": "--TEMP_INIT 0.1 --TEMP_FIM 5 --N_REPLICAS 11 --MCL 500 --PTL 500 --PASSO_GATILHO 10 "
                         "--TEMP_DIST 3 --TYPE_UPDATE 1 --INIT_SOL_TYPE 0 --TEMP_UPDATE 3500 "
                         "--PTL_TEMP_UPDATE_PROPORTION 3  (a planilha original dizia --PTL 600, mas o maior "
                         "PTL nos relatórios é 499)",
}

README = """\
Gerado por scripts/buildSpreadsheet.py a partir de output-final/ em {date}.
Valores: {mode_text}

ORIGEM DE CADA ABA (pastas em output-final/)
{sources}

COMO CADA COLUNA É CALCULADA
- PT e GA: 10 execuções por instância. Contadores (tarefas, trocas) = média das 10.
- PT: S* = maior FO das 10; S = média; σ = desvio-padrão amostral / média × 100;
  S0 = média de "Best Initial"; gap = (S* − S0) / S* × 100; tempo = média, em s;
  Convergência = média(PTL) / limite de PTL × 100 (limite 600 em Same/Different, 500 em Dang).
- GA: determinístico (semente 42): as 10 execuções dão a mesma FO. Resultado = média; Tempo = média, em s.
- Practitioner e MODELO: execução única. MODELO: Best Bound, Resultado e Tempo (s) do Gurobi.
- IPMTC II: FO Dang = FO com os pesos do Beezao (troca 1, tarefa não finalizada 30, lucro 0);
  FO Real = 30·finalizadas − trocas − 10·instâncias de troca − 30·não finalizadas (pesos da dissertação;
  toda operação do Beezao é prioritária). Tempo em s. PT: médias das 10 execuções.
- FO da dissertação = 30·finalizadas − trocas − 10·instâncias de troca − 30·prioritárias não finalizadas.
- Coluna Validação: "ok" ou as verificações que falharam (detalhes na aba Validação).

PARÂMETROS DO PT (copiados da planilha original; não ficam registrados nos relatórios)
{params}

PROBLEMAS ENCONTRADOS NA VALIDAÇÃO
- Practitioner (todas as abas): src/practitioner.cpp imprime em "unfineshedPriorityCount" o número de
  tarefas prioritárias FINALIZADAS, e a FO usa esse valor. No modo "recomputed" as colunas de prioritárias
  não finalizadas e Resultado/FO Dang são corrigidas (aba Correções).
- Different Toolset, PT e GA: o decodificador divide a sequência entre as máquinas pela quantidade de operações
  (splitSolutionIntoMachines) e pode separar a operação 0 e a 1 de uma tarefa reentrante (no máximo um par por
  fronteira entre máquinas), quebrando a cadeia forte; às vezes conta a operação 1 como finalizada sem a 0.
- Different Toolset, todos os métodos: loadData.cpp groupJobs soma o tempo da operação 1 na operação 0 sem
  agrupá-las, então a operação 1 é processada duas vezes em tempo (tempo desperdiçado; aviso duration-reentrant-sum).
- Different Toolset, Practitioner: não impõe a cadeia forte (operações separadas, fora de ordem, 1 sem 0).
- Different Toolset, MODELO: a restrição (9) do modelo é só precedência no tempo (cadeia fraca): operações 0 e 1
  em máquinas diferentes ou com outra operação entre elas.
- Same Toolset n=212 (v6, v7, v8): uma operação precisa de 105 ferramentas com capacidade 80 (nunca pode ser
  processada). PT e GA não a finalizam; o Practitioner a processa com 80 ferramentas e a conta como finalizada.
- Avisos que não mudam a viabilidade: pares reentrantes impressos com os tempos de op0/op1 trocados; MODELO
  imprime s/e truncados para inteiro (1 minuto de diferença); MODELO permite troca exatamente no minuto 720;
  célula "1z" em n=75,p=0.24 (lida como 1, como faz o std::stoi).

NÃO INCLUÍDO
- A tabela lateral do GA na aba "IPMTC II Pratictitioners e GA" da planilha original (931: 745, ...):
  nenhuma execução em output/ corresponde a ela.
"""


# ---------------------------------------------------------------------------------------------------
# writing
# ---------------------------------------------------------------------------------------------------

def _cell(v, digits=2):
    if isinstance(v, float):
        return int(v) if v.is_integer() else round(v, digits)
    return v


def write_xlsx(path, tabs, mode, corrections):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="center")
    title_font = Font(bold=True, size=13)
    bad_fill = PatternFill("solid", fgColor="F8D7DA")

    wb = Workbook()
    ws = wb.active
    ws.title = "LEIA-ME"
    sources = []
    for t in rs.TABS:
        for blk in tabs[t.tab]:
            sources.append(f"- {t.tab} / {blk.title}: {os.path.relpath(blk.source, REPO)}")
    text = README.format(
        date=datetime.date.today().isoformat(),
        mode_text=("recalculados a partir do sequenciamento impresso (validador)" if mode == "recomputed"
                   else "como impressos no rodapé dos relatórios (igual à planilha original)"),
        sources="\n".join(sources),
        params="\n".join(f"- {k}: {v}" for k, v in PT_PARAMS.items()))
    for i, line in enumerate(text.splitlines(), start=1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 130

    for tab_name, blocks in tabs.items():
        ws = wb.create_sheet(tab_name)
        r = 1
        widths = defaultdict(int)
        for blk in blocks:
            ws.cell(r, 1, blk.title).font = title_font
            r += 2
            for c, h in enumerate(blk.headers, start=1):
                cell = ws.cell(r, c, h)
                cell.font, cell.alignment = bold, wrap
            ws.row_dimensions[r].height = 45
            r += 1
            beezao = blk.headers[0] == "Instancia"  # Beezao times are a few ms: keep 4 decimals
            for row in blk.rows:
                for c, (v, h) in enumerate(zip(row, blk.headers), start=1):
                    cell = ws.cell(r, c, _cell(v, 4 if beezao and h == "Tempo" else 2))
                    widths[c] = max(widths[c], len(str(cell.value)))
                    if h == rs.VALIDATION_HEADER and v != "ok":
                        cell.fill = bad_fill
                r += 1
            r += 2
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 60)
        ws.freeze_panes = "A2"

    # validation detail
    ws = wb.create_sheet("Validação")
    head = ["Aba", "Bloco", "Instância", "Relatórios", "Relatórios com erro", "Erros", "Avisos"]
    ws.append(head)
    for tab_name, blocks in tabs.items():
        for blk in blocks:
            for key, results in zip(blk.keys, blk.status):
                err, warn = rs.issue_summary(results)
                if not err and not warn:
                    continue
                ws.append([tab_name, blk.title, _key_text(key), len(results), sum(not x.ok for x in results),
                           "; ".join(f"{k} ×{n}" for k, n in sorted(err.items())),
                           "; ".join(f"{k} ×{n}" for k, n in sorted(warn.items()))])
    ws.append([])
    ws.append(["Verificação", "Descrição"])
    for k, d in vd.CHECKS.items():
        ws.append([k, d])
    for c, w in zip("ABCDEFG", (22, 14, 44, 11, 11, 60, 70)):
        ws.column_dimensions[c].width = w
    for cell in ws[1]:
        cell.font = bold

    if corrections is not None:
        ws = wb.create_sheet("Correções")
        ws.append(["Aba", "Bloco", "Instância", "Coluna", "Valor impresso (reported)", "Valor recalculado"])
        for row in corrections:
            ws.append([row[0], row[1], _key_text(row[2]), row[3], _cell(row[4]), _cell(row[5])])
        for c, w in zip("ABCDEF", (22, 14, 44, 36, 24, 20)):
            ws.column_dimensions[c].width = w
        for cell in ws[1]:
            cell.font = bold

    wb.save(path)


def _key_text(key):
    if isinstance(key, tuple):
        return f"n={key[0]}, p={key[1]:g}, r={key[2]:g}"
    return str(key)


def diff_modes(tabs_rep, tabs_rec):
    """Cells whose value changes between reported and recomputed."""
    out = []
    for tab_name, blocks in tabs_rec.items():
        for b_rec, b_rep in zip(blocks, tabs_rep[tab_name]):
            for key, row_rec, row_rep in zip(b_rec.keys, b_rec.rows, b_rep.rows):
                for h, a, b in zip(b_rec.headers, row_rep, row_rec):
                    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) > 1e-9:
                        out.append((tab_name, b_rec.title, key, h, a, b))
    return out


# ---------------------------------------------------------------------------------------------------
# comparison with the Google Sheet CSV export
# ---------------------------------------------------------------------------------------------------

def _readings(s):
    """Every plausible number a Google Sheets pt-BR CSV cell can mean -> [(value, decimals)].
    The sheet mixes '14,53', '1.14', '1,059,50' (a thousands separator gone wrong) and '-'."""
    s = s.strip().replace("\xa0", "")
    if s in ("", "-"):
        return None
    try:
        if "," not in s:
            return [(float(s), len(s.split(".")[1]) if "." in s else 0)]
        p = s.split(",")
        if len(p) == 2:
            out = [(float(f"{p[0]}.{p[1]}"), len(p[1]))]
            if len(p[1]) == 3:
                out.append((float(p[0] + p[1]), 0))
            return out
        if len(p) == 3:
            return [(float(f"{p[0]}{p[1]}.{p[2]}"), len(p[2]))]
    except ValueError:
        pass
    return []


def _matches(mine, cell, tol=0.0):
    c = _readings(cell)
    if c is None or not isinstance(mine, (int, float)):
        return c is None and not isinstance(mine, (int, float))
    return any(abs(mine - v) <= max(0.5 * 10 ** (-d), tol) + 1e-9 for v, d in c)


# The sheet's Same/Different "Convergencia %" cells are text typed with two
# decimals ('1.14', '08.06') that differ from mean(PTL)/600 by up to 0.01.
COMPARE_TOL = {"Convergencia %": 0.01}


def _sheet_rows(path, grid):
    """-> {block title: {key: row}} of one exported tab."""
    rows = list(csv.reader(open(path, encoding="utf-8")))
    out, cur = defaultdict(dict), None
    for r in rows:
        filled = [c for c in r if c.strip()]
        if not filled:
            continue
        first = r[0].strip()
        if grid:
            if len(filled) == 1 and first in ("PT", "Practitioner", "MODELO", "GA"):
                cur = first
            elif cur and first.isdigit():
                key = (int(first), _readings(r[1])[0][0], _readings(r[2])[0][0])
                out[cur][key] = r
        elif first.startswith("instance"):
            out[None][first] = r
    return out


def compare(tabs, sheet_dir):
    total_bad = 0
    for t in rs.TABS:
        path = os.path.join(sheet_dir, f"Resultados SSP-USPrC - {SHEET_TAB.get(t.tab, t.tab)}.csv")
        if not os.path.exists(path):
            print(f"{t.tab}: {path} not found, skipped")
            continue
        grid = isinstance(t, rs.Grid)
        sheet = _sheet_rows(path, grid)
        for blk in tabs[t.tab]:
            srows = sheet.get(blk.title if grid else None, {})
            stats, bad = Counter(), defaultdict(list)
            mine_keys = set(blk.keys)
            for key, row in zip(blk.keys, blk.rows):
                srow = srows.get(key)
                if srow is None:
                    bad["<row missing in sheet>"].append(_key_text(key))
                    continue
                if grid and all(c.strip() in ("", "-") for c in srow[3:9]):
                    continue  # the sheet leaves unsolved instances as '-'
                for c, h in enumerate(blk.headers):
                    if h == rs.VALIDATION_HEADER or c >= len(srow) or row[c] == "-" and not srow[c].strip("-0 "):
                        continue
                    if c == 0 and not grid:
                        continue
                    ok = _matches(row[c], srow[c], COMPARE_TOL.get(h, 0.0))
                    stats[(h, ok)] += 1
                    if not ok:
                        bad[h].append(f"{_key_text(key)}: generated {row[c]!r}, sheet {srow[c]!r}")
            for k, srow in srows.items():
                if k not in mine_keys and not (grid and all(c.strip() in ("", "-") for c in srow[3:9])):
                    bad["<row missing in output-final>"].append(_key_text(k))
            cols = {h for h, _ in stats}
            n_bad = sum(stats[(h, False)] for h in cols) + sum(len(v) for k, v in bad.items() if k.startswith("<"))
            total_bad += n_bad
            print(f"{t.tab} / {blk.title}: {sum(stats.values())} cells compared, {n_bad} differences")
            for h, lines in bad.items():
                print(f"    {h}: {len(lines)}")
                for line in lines[:5]:
                    print(f"        {line}")
    return total_bad


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", default=os.path.join(REPO, "output-final"), help="folder with the final runs")
    ap.add_argument("-o", "--output", help="xlsx to write (default: <root>/Resultados-SSP-USPrC.xlsx)")
    ap.add_argument("--values", choices=["recomputed", "reported"], default="recomputed")
    ap.add_argument("--compare", metavar="DIR", help="compare with the Google Sheet tabs exported as CSV in DIR")
    ap.add_argument("--no-write", action="store_true", help="do not write the xlsx (e.g. only --compare)")
    ap.add_argument("--jobs", type=int, default=os.cpu_count())
    args = ap.parse_args()

    tabs, validated = rs.build(args.root, args.values, args.jobs)
    n_bad = sum(not r.ok for r in validated.values())
    print(f"{len(validated)} reports validated, {n_bad} with errors (python3 scripts/validateRuns.py for details)")

    if not args.no_write:
        corrections = None
        if args.values == "recomputed":
            tabs_rep, _ = rs.build(args.root, "reported", validated=validated)
            corrections = diff_modes(tabs_rep, tabs)
        out = args.output or os.path.join(args.root, "Resultados-SSP-USPrC.xlsx")
        write_xlsx(out, tabs, args.values, corrections)
        print(f"written {out} ({args.values} values"
              + (f", {len(corrections)} corrected cells)" if corrections is not None else ")"))
    if args.compare:
        print()
        sys.exit(1 if compare(tabs, args.compare) else 0)


if __name__ == "__main__":
    main()
